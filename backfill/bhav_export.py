"""Downloads of the bhavcopy check (dashboard Download menus, nifty_bhav.py --export / --export-days):
CSV and Excel files built from what a check already saved (bhavcopy/bhav_check.json, bhav_compare.csv
and summary.csv). Nothing is compared again and nothing in the day folder is changed.

Numbers are written raw (25000, 1571.5; never 20,36,45,585), dates as YYYY-MM-DD, and each CSV starts
with its header line, so Excel and pandas.read_csv read them as they are."""
import csv
import io
import os
import re

from . import bhav, fno

# A contract row: the dashboard's combined "bhav / Kite" columns split into one value per column
CONTRACT_COLUMNS = ["Contract", "Expiry", "Bucket", "Type", "Strike", "Bhav_volume", "Bhav_volume_shares", "Kite_volume",
                    "Bhav_OI", "Kite_OI", "Bhav_last", "Kite_last", "Kite_regular_close", "Kite_last_bar_time",
                    "Bhav_close", "Bhav_settle", "Status", "Status_note"]
# One row per checked day, as the day table
DAY_COLUMNS = ["Date", "Status", "Liquid_missing", "Liquid_diffs", "Bhav_traded", "Kite_files", "Missing", "Extra",
               "Volume_diffs", "Zero_volume_in_kite", "OI_diffs", "Price_diffs", "Toggle_keep_zero_volume"]
DAY_COUNTS = {"Liquid_missing": "liquid_missing", "Liquid_diffs": "liquid_diffs", "Bhav_traded": "bhav_traded",
              "Kite_files": "kite_files", "Missing": "missing", "Extra": "extra", "Volume_diffs": "volume_diffs",
              "Zero_volume_in_kite": "zero_volume", "OI_diffs": "oi_diffs", "Price_diffs": "price_diffs"}
EXPIRY_COLUMNS = ["Expiry", "Bucket", "Listed_strikes", "Traded_strikes", "Bhav_traded_contracts", "Kite_files",
                  "Missing", "Extra", "Volume_diffs", "OI_diffs", "Price_diffs"]
SETTLEMENT_COLUMNS = ["Contract", "Expiry", "Type", "Strike", "Bhav_close", "Bhav_settle", "Bhav_OI"]
SHEETS = ["Summary", "Not OK", "All contracts", "Settlement"]
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
NUMBER = re.compile(r"-?\d+(\.\d+)?")


# --- rows -------------------------------------------------------------------------------------

def contract_row(r, unit):
    """A bhav_compare.csv row in CONTRACT_COLUMNS (older files used Kite_last_any / Kite_last_1530)."""
    vol = r.get("Bhav_volume", "")
    return {"Contract": r["Contract"], "Expiry": r.get("Expiry", ""), "Bucket": r.get("Bucket", bhav.FAR),
            "Type": r.get("Type", ""), "Strike": r.get("Strike", ""), "Bhav_volume": vol,
            "Bhav_volume_shares": fno.fmt_num(bhav.num(vol) * unit) if vol != "" else "",
            "Kite_volume": r.get("Kite_volume", ""), "Bhav_OI": r.get("Bhav_OI", ""), "Kite_OI": r.get("Kite_OI", ""),
            "Bhav_last": r.get("Bhav_last", ""), "Kite_last": r.get("Kite_last") or r.get("Kite_last_any", ""),
            "Kite_regular_close": r.get("Kite_last_regular") or r.get("Kite_last_1530", ""),
            "Kite_last_bar_time": r.get("Last_bar", ""), "Bhav_close": r.get("Bhav_close", ""),
            "Bhav_settle": r.get("Bhav_settle", ""), "Status": r["Status"], "Status_note": r.get("Note", "")}


def relabel(columns, rows, source="kite"):
    """Kite files keep the Kite_* names. For Dhan files (source="dhan") the same tables say Dhan:
    Kite_volume -> Dhan_volume, MISSING_IN_KITE -> MISSING_IN_DHAN, "Kite's last bar" -> "Dhan's last bar"."""
    if source == "kite":
        return columns, rows
    name = fno.BROKER[source]
    swap = lambda t: t.replace("Kite", name).replace("kite", name.lower()).replace("KITE", name.upper())
    return [swap(c) for c in columns], [{swap(k): swap(v) if isinstance(v, str) else v for k, v in r.items()} for r in rows]


def volume_unit(check):
    return (check.get("units") or {}).get("volume", 1)


def all_contracts(data_root, underlying, run_date, source="kite"):
    """(every contract of the day in CONTRACT_COLUMNS, check). Raises BhavError when not checked."""
    rows, check = bhav.compare_rows(data_root, underlying, run_date, source)
    unit = volume_unit(check)
    return [contract_row(r, unit) for r in rows], check


def view_rows(data_root, underlying, run_date, bucket="all", source="kite"):
    """The rows the dashboard's difference list shows for a bucket filter ("all", "liquid" or a bucket
    name), in its order (biggest bhavcopy volume first)."""
    rows, check = bhav.diff_rows(data_root, underlying, run_date, source)
    if bucket == "liquid":
        rows = [r for r in rows if r["Liquid"]]
    elif bucket != "all":
        rows = [r for r in rows if r["Bucket"] == bucket]
    unit = volume_unit(check)
    return [contract_row(r, unit) for r in rows]


def day_status(d):
    if d["status"] == "ok":
        return "OK"
    return "CHECK (minor)" if d["severity"] == "minor" else "CHECK"


def day_row(data_root, underlying, d, source="kite"):
    """A list_days() entry in DAY_COLUMNS; counts an older check doesn't have stay blank."""
    counts = d["counts"]
    day_dir = os.path.join(fno.fno_root(data_root, underlying, source), d["date"])
    return {"Date": d["date"], "Status": day_status(d),
            **{col: counts.get(key, "") for col, key in DAY_COUNTS.items()},
            "Toggle_keep_zero_volume": fno.saved_keep_zero(day_dir)}


def days_rows(data_root, underlying, source="kite"):
    """One row per checked day, latest first. Raises BhavError when no day has been checked."""
    out = [day_row(data_root, underlying, d, source) for d in bhav.list_days(data_root, underlying, source)
           if d["status"] != "pending"]
    if not out:
        raise bhav.BhavError("No day in this folder has been checked against the bhavcopy yet.")
    return out


def expiry_rows(data_root, underlying, run_date, contracts, source="kite"):
    """Per-expiry rows from summary.csv; Volume_diffs counted from the day's contract rows."""
    path = os.path.join(fno.fno_root(data_root, underlying, source), run_date.isoformat(), "summary.csv")
    if not os.path.exists(path):
        return []
    volume = {}
    for r in contracts:
        if "VOLUME_DIFF" in r["Status"]:
            key = (r["Expiry"], "FUT" if r["Type"] == "FUT" else "OPT")
            volume[key] = volume.get(key, 0) + 1
    with open(path, newline="", encoding="utf-8") as f:
        summary = list(csv.DictReader(f))
    checked = any(s.get("Bhav_traded_contracts") not in (None, "") for s in summary)
    return [{"Expiry": s.get("Expiry", ""), "Bucket": s.get("Bucket", ""), "Listed_strikes": s.get("Listed_strikes", ""),
             "Traded_strikes": s.get("Traded_strikes", ""), "Bhav_traded_contracts": s.get("Bhav_traded_contracts", ""),
             "Kite_files": s.get("Kite_files", ""), "Missing": s.get("Missing_in_kite", ""),
             "Extra": s.get("Extra_in_kite", ""),
             "Volume_diffs": volume.get((s.get("Expiry", ""), "FUT" if s.get("Bucket", "").startswith("Futures") else "OPT"), 0)
             if checked else "",
             "OI_diffs": s.get("OI_diffs", ""), "Price_diffs": s.get("Price_diffs", "")} for s in summary]


# --- files ------------------------------------------------------------------------------------

def csv_bytes(columns, rows, source="kite"):
    columns, rows = relabel(columns, rows, source)
    buf = io.StringIO()
    w = csv.DictWriter(buf, fieldnames=columns, lineterminator="\n", extrasaction="ignore")
    w.writeheader()
    w.writerows(rows)
    return buf.getvalue().encode("utf-8")


def cell(v):
    """Numbers as numbers in Excel; everything else (names, dates, times) as text."""
    if isinstance(v, str) and NUMBER.fullmatch(v):
        return float(v) if "." in v else int(v)
    return v


def put_table(ws, top, columns, rows, source="kite"):
    """Header row (bold) at row top, then the rows; returns the last row number."""
    from openpyxl.styles import Font
    columns, rows = relabel(columns, rows, source)
    for c, name in enumerate(columns, 1):
        ws.cell(top, c, name).font = Font(bold=True)
    for i, r in enumerate(rows, top + 1):
        for c, name in enumerate(columns, 1):
            ws.cell(i, c, cell(r.get(name, "")))
    return top + len(rows)


def fit_columns(ws):
    from openpyxl.utils import get_column_letter
    widths = {}
    for row in ws.iter_rows():
        for c in row:
            if c.value is not None:
                widths[c.column] = max(widths.get(c.column, 0), len(str(c.value)))
    for col, w in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = min(w + 2, 60)


def table_sheet(wb, title, columns, rows, source="kite"):
    from openpyxl.utils import get_column_letter
    ws = wb.create_sheet(title)
    last = put_table(ws, 1, columns, rows, source)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{max(last, 1)}"
    fit_columns(ws)


def xlsx_bytes(data_root, underlying, run_date, source="kite"):
    """The full report of a day: Summary (the day row, then one row per expiry), Not OK, All contracts
    and Settlement (bhavcopy close, settlement price and OI per contract, for backtest P&L)."""
    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter
    contracts, check = all_contracts(data_root, underlying, run_date, source)
    day = next((d for d in bhav.list_days(data_root, underlying, source) if d["date"] == run_date.isoformat()), None)
    wb = Workbook()
    ws = wb.active
    ws.title = SHEETS[0]
    put_table(ws, 1, DAY_COLUMNS, [day_row(data_root, underlying, day, source)] if day else [], source)
    expiries = expiry_rows(data_root, underlying, run_date, contracts, source)
    last = put_table(ws, 4, EXPIRY_COLUMNS, expiries, source)
    ws.freeze_panes = "A5"   # the day row and the expiry header stay in view
    ws.auto_filter.ref = f"A4:{get_column_letter(len(EXPIRY_COLUMNS))}{max(last, 4)}"
    fit_columns(ws)
    table_sheet(wb, SHEETS[1], CONTRACT_COLUMNS, [r for r in contracts if r["Status"] != "OK"], source)
    table_sheet(wb, SHEETS[2], CONTRACT_COLUMNS, contracts, source)
    table_sheet(wb, SHEETS[3], SETTLEMENT_COLUMNS, [r for r in contracts if r["Bhav_settle"] != ""], source)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def slug(text):
    return re.sub(r"[^a-z0-9]+", "_", text.lower()).strip("_") or "all"


def export(data_root, underlying, scope, run_date=None, bucket="all", today=None, source="kite"):
    """(file name, bytes, media type) for scope "days", "view", "all" or "xlsx". Reads only.
    Raises BhavError when the day (or, for "days", every day) has no check. Dhan files (source="dhan")
    get Dhan_* columns and a _dhan file name suffix; underlyings other than NIFTY add theirs (_sensex)."""
    tag = ("" if underlying == "NIFTY" else f"_{underlying.lower()}") + ("" if source == "kite" else f"_{source}")
    if scope == "days":
        today = today or fno.today_ist()
        return (f"bhavcopy_check_all_days_{today}{tag}.csv",
                csv_bytes(DAY_COLUMNS, days_rows(data_root, underlying, source), source), "text/csv")
    if scope == "view":
        return (f"bhavcopy_check_{run_date}_{slug(bucket)}{tag}.csv",
                csv_bytes(CONTRACT_COLUMNS, view_rows(data_root, underlying, run_date, bucket, source), source), "text/csv")
    if scope == "all":
        return (f"bhavcopy_check_{run_date}_all_contracts{tag}.csv",
                csv_bytes(CONTRACT_COLUMNS, all_contracts(data_root, underlying, run_date, source)[0], source), "text/csv")
    if scope == "xlsx":
        return f"bhavcopy_check_{run_date}{tag}.xlsx", xlsx_bytes(data_root, underlying, run_date, source), XLSX
    raise ValueError(f"unknown export scope {scope!r}")
