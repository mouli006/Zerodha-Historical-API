"""NSE F&O bhavcopy check (backfill/bhav.py, nifty_bhav.py, dashboard routes) on synthetic UDiFF files."""
import csv
import io
import json
import os
import zipfile
from datetime import date, datetime, timedelta

import pandas as pd
import pytest
import requests

import dashboard
import nifty_bhav
from app import app
from backfill import bhav, bhav_export, fno
from backfill.job import JOB, JOB_LOCK
from backfill.storage import open_db
from test_fno import RUN_DATE, ACCESS, FakeKite, bars, fake_kite, read_csv, run, set_now  # noqa: F401

IST = fno.IST
DAY = "2026-10-12"


def row(tp, expiry, strike="", optn="", vol=0, oi=0, last=0, close=0, settle=0, sym="NIFTY", lot=50, day=DAY):
    r = dict.fromkeys(bhav.BHAV_HEADER, "")
    r.update(TradDt=day, BizDt=day, Sgmt="FO", Src="NSE", FinInstrmTp=tp, TckrSymb=sym, XpryDt=expiry,
             FininstrmActlXpryDt=expiry, StrkPric=strike, OptnTp=optn, LastPric=last, ClsPric=close,
             SttlmPric=settle, OpnIntrst=oi, TtlTradgVol=vol, NewBrdLotQty=lot, SsnId="F1")
    return r


def csv_bytes(rows):
    buf = io.StringIO()
    w = csv.DictWriter(buf, fieldnames=bhav.BHAV_HEADER, lineterminator="\n")
    w.writeheader()
    w.writerows(rows)
    return buf.getvalue().encode()


def zip_bytes(rows, day=DAY):
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr(f"BhavCopy_NSE_FO_0_0_0_{day.replace('-', '')}_F_0000.csv", csv_bytes(rows))
    return buf.getvalue()


def kite_file(day_dir, name, closes, vol=100, oi_last=1000, first="09:15"):
    """A saved day file with one bar per close; the last bar carries oi_last."""
    h, m = map(int, first.split(":"))
    t0 = datetime(2026, 10, 12, h, m, tzinfo=IST)
    candles = [{"date": t0 + timedelta(minutes=k), "open": c, "high": c, "low": c, "close": c,
                "volume": vol, "oi": oi_last} for k, c in enumerate(closes)]
    sub = "options" if name.endswith(("CE", "PE")) else "futures"
    fno.write_contract(os.path.join(day_dir, sub, name + ".csv"), name, candles)


@pytest.fixture(autouse=True)
def no_network(monkeypatch):
    class Blocked:
        def __init__(self, *a, **kw):
            raise AssertionError("network used")
    monkeypatch.setattr(bhav.requests, "Session", Blocked)


# --- parsing ----------------------------------------------------------------------------------

def test_name_mapping():
    items = bhav.parse_bhav([row("IDO", "2026-10-13", "22650.00", "PE", vol=4),
                             row("IDF", "2026-10-27", vol=9),
                             row("IDF", "2026-10-27", sym="BANKNIFTY", vol=9),
                             row("STF", "2026-10-27", sym="NIFTY", vol=9)])
    assert [(i["name"], i["type"], i["traded"]) for i in items] == [
        ("NIFTY2265013OCT26PE", "PE", True), ("NIFTY27OCT26", "FUT", True)]


def test_bad_strike_is_reported_and_check_goes_on(tmp_path):
    items = bhav.parse_bhav([row("IDO", "2026-10-13", "22650.50", "CE", vol=1), row("IDF", "2026-10-27", vol=1)])
    assert items[0]["error"] and items[1]["name"] == "NIFTY27OCT26"
    rows, counts, _ = bhav.compare(items, str(tmp_path))
    assert counts["bad_names"] == 1 and rows[0]["Status"].startswith("BAD_NAME")


def test_read_bhav_zip_and_csv():
    rows = [row("IDF", "2026-10-27", vol=1)]
    assert bhav.read_bhav(zip_bytes(rows))[1][0]["XpryDt"] == "2026-10-27"
    assert bhav.read_bhav(csv_bytes(rows))[1][0]["XpryDt"] == "2026-10-27"
    with pytest.raises(bhav.BhavError):
        bhav.read_bhav(b"a,b\n1,2\n")


# --- comparing --------------------------------------------------------------------------------

def statuses(rows):
    return {r["Contract"]: r["Status"] for r in rows}


def test_coverage_missing_and_extra(tmp_path):
    d = str(tmp_path)
    kite_file(d, "NIFTY27OCT26", [100.0], vol=650)                 # A: future
    kite_file(d, "NIFTY2500013OCT26CE", [10.0], vol=650)           # B
    kite_file(d, "NIFTY2510013OCT26CE", [5.0], vol=650)            # D: on disk, volume 0 in bhav
    items = bhav.parse_bhav([
        row("IDF", "2026-10-27", vol=10, last=100, lot=65, oi=1000),
        row("IDO", "2026-10-13", "25000.00", "CE", vol=10, last=10, lot=65, oi=1000),
        row("IDO", "2026-10-13", "25050.00", "CE", vol=3, last=7, lot=65),        # C: traded, no file
        row("IDO", "2026-10-13", "25100.00", "CE", vol=0, lot=65),
    ])
    rows, counts, units = bhav.compare(items, d)
    st = statuses(rows)
    assert st["NIFTY2505013OCT26CE"] == "MISSING_IN_KITE"
    assert st["NIFTY2510013OCT26CE"] == "EXTRA_IN_KITE"
    assert st["NIFTY27OCT26"] == "OK" and st["NIFTY2500013OCT26CE"] == "OK"
    assert (counts["missing"], counts["extra"], counts["bhav_traded"], counts["kite_files"]) == (1, 1, 3, 3)
    assert bhav.check_status(counts) == "Bhavcopy check: CHECK (1 missing, 1 extra); liquid buckets: 1 missing, 1 diffs"


def test_units_detected_and_diffs_flagged(tmp_path):
    # bhav volume in contracts (lot 65), OI in shares; Kite has shares for both
    d = str(tmp_path)
    kite_file(d, "NIFTY27OCT26", [100.0] * 10, vol=6500, oi_last=130000)       # 65,000 shares = 1,000 lots
    kite_file(d, "NIFTY24NOV26", [101.0] * 4, vol=6500, oi_last=65000)         # 26,000 shares = 400 lots
    kite_file(d, "NIFTY2500013OCT26CE", [10.0], vol=66950, oi_last=20800)      # +3% = 30 lots; OI +4%
    kite_file(d, "NIFTY2500013OCT26PE", [12.0], vol=6695, oi_last=20500)       # +3% but only 3 lots; OI +2.5%
    kite_file(d, "NIFTY2510013OCT26PE", [12.0], vol=65325, oi_last=20000)      # +0.5% = 5 lots: under 1%
    items = bhav.parse_bhav([
        row("IDF", "2026-10-27", vol=1000, oi=130000, last=100, lot=65),
        row("IDF", "2026-11-24", vol=400, oi=65000, last=101, lot=65),
        row("IDO", "2026-10-13", "25000.00", "CE", vol=1000, oi=20000, last=10, lot=65),
        row("IDO", "2026-10-13", "25000.00", "PE", vol=100, oi=20000, last=12, lot=65),
        row("IDO", "2026-10-13", "25100.00", "PE", vol=1000, oi=20000, last=12, lot=65),
    ])
    rows, counts, units = bhav.compare(items, d)
    assert units == {"volume": 65, "oi": 1}
    st = statuses(rows)
    assert st["NIFTY2500013OCT26CE"] == "VOLUME_DIFF+OI_DIFF"
    assert st["NIFTY2500013OCT26PE"] == "OK"        # 3% off but within 5 lots; OI 2.5% is within 3%
    assert st["NIFTY2510013OCT26PE"] == "OK" and st["NIFTY27OCT26"] == "OK"
    assert (counts["volume_diffs"], counts["oi_diffs"]) == (1, 1)


def test_price_against_last_bar(tmp_path):
    d = str(tmp_path)
    # bars 15:29 (close 99), 15:30, 15:31, 15:32 (close 110): NSE's last price 110 matches the last bar;
    # the regular close is the 15:29 candle (Kite labels candles by their start time)
    kite_file(d, "NIFTY2500013OCT26CE", [99.0, 100.0, 108.0, 110.0], first="15:29", vol=50)
    kite_file(d, "NIFTY2500013OCT26PE", [50.0, 51.0], first="15:28", vol=50)           # last bar 15:29
    kite_file(d, "NIFTY2510013OCT26PE", [50.0] * 13, first="15:29", vol=50)            # last bar 15:41
    kite_file(d, "NIFTY27OCT26", [22600.0], vol=50)
    items = bhav.parse_bhav([
        row("IDO", "2026-10-13", "25000.00", "CE", vol=4, last=100.4, oi=1000),       # vs 110: differs
        row("IDO", "2026-10-13", "25000.00", "PE", vol=2, last=60, oi=1000),
        row("IDO", "2026-10-13", "25100.00", "PE", vol=13, last=55, oi=1000),
        row("IDF", "2026-10-27", vol=1, last=22600, close=22750, oi=1000),             # ClsPric far off: ignored
    ])
    rows, counts, _ = bhav.compare(items, d)
    by = {r["Contract"]: r for r in rows}
    ce = by["NIFTY2500013OCT26CE"]
    assert (ce["Kite_last"], ce["Kite_last_regular"], ce["Last_bar"]) == ("110", "99", "15:32:00")
    assert ce["Status"] == "PRICE_DIFF" and ce["Note"] == "trades after Kite's last bar"
    pe = by["NIFTY2500013OCT26PE"]
    assert (pe["Status"], pe["Last_bar"], pe["Note"]) == ("PRICE_DIFF", "15:29:00", "trades after Kite's last bar")
    late = by["NIFTY2510013OCT26PE"]
    assert (late["Status"], late["Last_bar"], late["Note"]) == ("PRICE_DIFF", "15:41:00", "")   # no excuse after 15:40
    fut = by["NIFTY27OCT26"]
    assert fut["Status"] == "OK" and fut["Bhav_close"] == "22750" and fut["Note"] == ""
    assert counts["price_diffs"] == 3


def test_price_within_tolerance(tmp_path):
    d = str(tmp_path)
    kite_file(d, "NIFTY2500013OCT26CE", [10.0, 10.1], first="15:29", vol=50)          # 2 ticks from 10.0
    kite_file(d, "NIFTY27OCT26", [22600.0, 22700.0], first="15:29", vol=50)           # 0.44% off
    items = bhav.parse_bhav([row("IDO", "2026-10-13", "25000.00", "CE", vol=2, last=10.0, oi=1000),
                             row("IDF", "2026-10-27", vol=2, last=22600, oi=1000)])
    assert set(statuses(bhav.compare(items, d)[0]).values()) == {"OK"}


def test_detect_unit():
    assert bhav.detect_unit([(6500, 100), (2600, 40)], 65, 65) == 65
    assert bhav.detect_unit([(1000, 1000), (990, 1000)], 65, 1) == 1
    assert bhav.detect_unit([], 65, 65) == 65


# --- download -----------------------------------------------------------------------------------

class FakeResponse:
    def __init__(self, status, content=b""):
        self.status_code, self.content, self.ok = status, content, 200 <= status < 300


class FakeSession:
    def __init__(self, responses):
        self.responses, self.headers, self.urls = list(responses), {}, []

    def get(self, url, timeout=None):
        self.urls.append(url)
        if "www.nseindia.com" in url:
            return FakeResponse(403)          # the home page blocks scripts; ignored
        r = self.responses.pop(0)
        if isinstance(r, Exception):
            raise r
        return r


def test_fetch_ok_despite_home_page_403():
    s = FakeSession([FakeResponse(200, zip_bytes([row("IDF", "2026-10-27")]))])
    data, url = bhav.fetch_bhav(date(2026, 10, 12), log=lambda m: None, session=s, sleep=lambda x: None)
    assert data[:2] == b"PK" and url.endswith("BhavCopy_NSE_FO_0_0_0_20261012_F_0000.csv.zip")
    assert s.headers["Referer"] == "https://www.nseindia.com/"


def test_fetch_404_is_not_available():
    with pytest.raises(bhav.BhavNotAvailable):
        bhav.fetch_bhav(date(2026, 10, 12), log=lambda m: None, session=FakeSession([FakeResponse(404)]),
                        sleep=lambda x: None)


def test_fetch_retries_then_fails():
    waits = []
    s = FakeSession([requests.ConnectionError("down"), FakeResponse(503), FakeResponse(200, b"<html>")])
    with pytest.raises(bhav.BhavError, match="--bhav-file"):
        bhav.fetch_bhav(date(2026, 10, 12), log=lambda m: None, session=s, sleep=waits.append)
    assert waits == [2, 4]


def test_not_available_writes_only_a_log_line(tmp_path, monkeypatch):
    def fail(*a, **kw):
        raise bhav.BhavNotAvailable("404")
    monkeypatch.setattr(bhav, "fetch_bhav", fail)
    assert bhav.run_bhav(date(2026, 10, 12), str(tmp_path), log=lambda m: None) == ("not_available", bhav.PENDING, None)
    root = tmp_path / "NIFTY_FnO_1min"
    assert os.listdir(root) == ["bhavcopy.log"] and "2026-10-12" in (root / "bhavcopy.log").read_text()


@pytest.mark.parametrize("kind", ["zip", "csv"])
def test_bhav_file_without_network(tmp_path, kind):
    rows = [row("IDF", "2026-10-27", vol=3, last=22600), row("IDO", "2026-10-13", "25000.00", "CE", vol=2, last=10),
            row("IDF", "2026-10-27", sym="BANKNIFTY", vol=5)]
    src = tmp_path / f"manual.{kind}"
    src.write_bytes(zip_bytes(rows) if kind == "zip" else csv_bytes(rows))
    status, line, summary = bhav.run_bhav(date(2026, 10, 12), str(tmp_path / "data"), bhav_file=str(src), log=lambda m: None)
    out = tmp_path / "data" / "NIFTY_FnO_1min" / DAY / "bhavcopy"
    assert status == "diffs" and line == "Bhavcopy check: CHECK (2 missing); liquid buckets: 2 missing, 0 diffs" and summary is None
    assert (out / f"manual.{kind}").read_bytes() == src.read_bytes()        # original kept unchanged
    assert len(read_csv(out / "nifty_bhav.csv")) == 3                        # header + 2 NIFTY rows
    assert {r[0] for r in read_csv(out / "bhav_compare.csv")[1:]} == {"NIFTY27OCT26", "NIFTY2500013OCT26CE"}
    assert json.loads((out / "bhav_check.json").read_text())["counts"]["missing"] == 2


def test_wrong_day_file_refused(tmp_path):
    src = tmp_path / "other.zip"
    src.write_bytes(zip_bytes([row("IDF", "2026-10-27", day="2026-10-09")], day="2026-10-09"))
    with pytest.raises(bhav.BhavError, match="2026-10-09"):
        bhav.run_bhav(date(2026, 10, 12), str(tmp_path), bhav_file=str(src), log=lambda m: None)


# --- summary.csv and run.log of a real (fake Kite) day ---------------------------------------

def day_bhav(path):
    """Bhavcopy for the fake Kite day of test_fno (lot 50, so Kite share volumes divide exactly)."""
    rows = [
        row("IDF", "2026-10-27", vol=10, oi=1001, last=25122, close=25130),     # 5 bars x 100
        row("IDF", "2026-11-24", vol=8, oi=1003, last=25303),
        row("IDO", "2026-10-13", "25100.00", "CE", vol=10, oi=1004, last=124),
        row("IDO", "2026-10-13", "25100.00", "PE", vol=6, oi=1002, last=97),
        row("IDO", "2026-10-13", "25150.00", "CE", vol=3, oi=500, last=80),       # traded: Kite had no volume
        row("IDO", "2026-10-13", "25150.00", "PE", vol=0),
        row("IDO", "2026-10-20", "25100.00", "CE", vol=4, oi=1001, last=151),
        row("IDO", "2026-10-27", "25100.00", "PE", vol=0),
    ]
    path.write_bytes(zip_bytes(rows))
    return str(path)


def summary_by(path):
    rows = read_csv(path)
    return {(r[0], r[1]): dict(zip(rows[0], r)) for r in rows[1:]}


def test_summary_gets_bhav_columns_and_keeps_them(tmp_path, fake_kite):
    data = tmp_path / "data"
    run(data)
    day = data / "NIFTY_FnO_1min" / DAY
    assert summary_by(day / "summary.csv")[("2026-10-13", "Current week")]["Bhav_traded_contracts"] == ""
    status, line, rows = bhav.run_bhav(RUN_DATE, str(data), bhav_file=day_bhav(tmp_path / "b.zip"), log=lambda m: None)
    assert (status, line) == ("diffs", "Bhavcopy check: CHECK (1 missing); liquid buckets: 1 missing, 0 diffs")
    s = summary_by(day / "summary.csv")
    w1 = s[("2026-10-13", "Current week")]
    assert (w1["Bhav_traded_contracts"], w1["Kite_files"], w1["Missing_in_kite"], w1["Extra_in_kite"],
            w1["OI_diffs"], w1["Price_diffs"]) == ("3", "2", "1", "0", "0", "0")
    fut = s[("2026-10-27", "Futures")]
    assert (fut["Bhav_traded_contracts"], fut["Kite_files"], fut["Missing_in_kite"]) == ("1", "1", "0")
    assert s[("2026-10-27", "Current month")]["Bhav_traded_contracts"] == "0"
    assert "Bhavcopy check: CHECK (1 missing); liquid buckets: 1 missing, 0 diffs" in (day / "run.log").read_text(encoding="utf-8")
    assert [r for r in rows if r["Bucket"] == "Futures"][0]["Kite_files"] == 1
    # a Kite re-run of the day keeps the check
    run(data)
    assert summary_by(day / "summary.csv")[("2026-10-13", "Current week")]["Missing_in_kite"] == "1"
    assert "Bhavcopy check: CHECK (1 missing)" in JOB["message"]


# --- bhav_summary.csv, the difference files and the check after every Kite day ------------------

def dict_rows(path):
    with open(path, newline="", encoding="utf-8") as f:
        return list(csv.DictReader(f))


def test_check_writes_only_the_summary_files(tmp_path, fake_kite):
    data = tmp_path / "data"
    run(data)
    bhav.run_bhav(RUN_DATE, str(data), bhav_file=day_bhav(tmp_path / "b.zip"), log=lambda m: None)
    out = data / "NIFTY_FnO_1min" / DAY / "bhavcopy"
    summary = dict_rows(out / bhav.SUMMARY_FILE)
    top = summary[0]
    assert (top["Expiry"], top["Date"], top["Status"], top["Missing"], top["Bhav_traded"]) == ("ALL", DAY, "CHECK", "1", "6")
    assert (top["Volume_unit"], top["OI_unit"]) == ("50", "1")
    w1 = [r for r in summary if (r["Expiry"], r["Type"]) == ("2026-10-13", "OPT")][0]
    assert (w1["Bucket"], w1["Bhav_contracts"], w1["Bhav_traded"], w1["Missing"]) == ("Current week", "4", "3", "1")
    assert sum(int(r["Bhav_contracts"]) for r in summary[1:]) == int(top["Bhav_contracts"])
    assert [r["Contract"] for r in dict_rows(out / "bhav_compare.csv") if r["Status"] == "MISSING_IN_KITE"] == \
        ["NIFTY2515013OCT26CE"]
    # no per-difference files: the zip, the NIFTY rows, the comparison, the result and the two summaries only
    assert sorted(os.listdir(out)) == sorted(["b.zip", "nifty_bhav.csv", "bhav_compare.csv", "bhav_check.json",
                                              bhav.SUMMARY_FILE, bhav.DAYS_SUMMARY_FILE])


@pytest.fixture
def auto_bhav(monkeypatch, tmp_path):
    """Kite runs fetch the bhavcopy (tests/conftest.py turns that off); NSE answers from calls["data"]."""
    monkeypatch.setattr(fno, "AUTO_BHAV", True)
    calls = {"data": None, "fetched": []}

    def fetch(run_date, log=print, budget=None, **kw):
        calls["fetched"].append((run_date, budget))
        if calls["data"] is None:
            raise bhav.BhavNotAvailable(f"NSE has no F&O bhavcopy for {run_date} yet (404: test).")
        return calls["data"], bhav.BHAV_URL.format(d=f"{run_date:%Y%m%d}")
    monkeypatch.setattr(bhav, "fetch_bhav", fetch)
    return calls


def test_kite_run_saves_and_checks_the_bhavcopy(tmp_path, fake_kite, auto_bhav):
    auto_bhav["data"] = open(day_bhav(tmp_path / "b.zip"), "rb").read()
    data = tmp_path / "data"
    assert run(data) == "finished_failures"            # the bhavcopy never changes the download's status
    out = data / "NIFTY_FnO_1min" / DAY / "bhavcopy"
    assert auto_bhav["fetched"] == [(RUN_DATE, bhav.CHECK_BUDGET)]
    assert (out / "BhavCopy_NSE_FO_0_0_0_20261012_F_0000.csv.zip").exists() and (out / "nifty_bhav.csv").exists()
    for name in ["bhav_compare.csv", bhav.SUMMARY_FILE, bhav.DAYS_SUMMARY_FILE]:
        assert (out / name).exists(), name
    assert not (out / "missing_in_kite.csv").exists()
    check = json.loads((out / "bhav_check.json").read_text())
    assert check["counts"]["missing"] == 1 and check["source"].endswith("20261012_F_0000.csv.zip")
    assert summary_by(out.parent / "summary.csv")[("2026-10-13", "Current week")]["Missing_in_kite"] == "1"
    assert "Bhavcopy check: CHECK (1 missing)" in JOB["message"]
    assert bhav.list_days(str(data))[0]["status"] == "diffs"
    # a re-run uses the saved bhavcopy: NSE is not asked again
    run(data)
    assert len(auto_bhav["fetched"]) == 1


def test_kite_run_bhavcopy_not_published(tmp_path, fake_kite, auto_bhav):
    data = tmp_path / "data"
    assert run(data) == "finished_failures"
    root = data / "NIFTY_FnO_1min"
    assert not (root / DAY / "bhavcopy" / "bhav_check.json").exists()
    assert "2026-10-12: NSE has no F&O bhavcopy" in (root / "bhavcopy.log").read_text(encoding="utf-8")
    assert "Bhavcopy for 2026-10-12 not saved" in (root / DAY / "run.log").read_text(encoding="utf-8")
    assert "Bhavcopy: pending (NSE has no F&O bhavcopy" in JOB["message"]
    assert bhav.list_days(str(data))[0]["status"] == "pending"


def test_refetch_does_not_fetch_the_bhavcopy(tmp_path, fake_kite, auto_bhav, monkeypatch):
    data = tmp_path / "data"
    monkeypatch.setattr(fno, "AUTO_BHAV", False)
    run(data)
    monkeypatch.setattr(fno, "AUTO_BHAV", True)
    assert job_begin_fno(data)
    fno.refetch_no_trade("key", ACCESS, RUN_DATE, str(data), 1000, log=lambda m: None)
    assert auto_bhav["fetched"] == []


def job_begin_fno(data):
    from backfill.job import job_begin
    return job_begin(**fno.job_fields("NIFTY", RUN_DATE, str(data), 1000, keep_zero=True))


# --- CLI ----------------------------------------------------------------------------------------

def make_day(root, d, checked=False):
    day = root / "NIFTY_FnO_1min" / d
    day.mkdir(parents=True)
    (day / "summary.csv").write_text("x")
    if checked:
        (day / "bhavcopy").mkdir()
        (day / "bhavcopy" / "bhav_check.json").write_text("{}")


@pytest.fixture
def cli_runs(monkeypatch):
    calls = []

    def fake(d, folder, underlying, bhav_file, log=print, source="kite"):
        calls.append(d if source == "kite" else (d, source))
        return results.get(d, "ok"), "", None
    results = {}
    monkeypatch.setattr(nifty_bhav, "run_bhav", fake)
    return calls, results


def test_cli_date(cli_runs, tmp_path):
    calls, results = cli_runs
    assert nifty_bhav.main(["2026-10-12", "--folder", str(tmp_path)]) == 0
    results[date(2026, 10, 12)] = "diffs"
    assert nifty_bhav.main(["2026-10-12", "--folder", str(tmp_path)]) == 1
    results[date(2026, 10, 12)] = "not_available"
    assert nifty_bhav.main(["2026-10-12", "--folder", str(tmp_path)]) == 4


def test_cli_missing_picks_unchecked_days(cli_runs, tmp_path):
    calls, results = cli_runs
    make_day(tmp_path, "2026-10-07")
    make_day(tmp_path, "2026-10-08", checked=True)
    make_day(tmp_path, "2026-10-09")
    results[date(2026, 10, 9)] = "not_available"
    assert nifty_bhav.main(["--missing", "--folder", str(tmp_path)]) == 4
    assert calls == [date(2026, 10, 7), date(2026, 10, 9)]


@pytest.mark.parametrize("argv", [[], ["2026-10-12", "--missing"], ["--missing", "--bhav-file", "x.zip"], ["12-10-2026"]])
def test_cli_needs_one_date_or_missing(cli_runs, tmp_path, argv):
    with pytest.raises(SystemExit) as e:
        nifty_bhav.main(argv + ["--folder", str(tmp_path)])
    assert e.value.code == 1 and cli_runs[0] == []


def test_cli_shares_the_kite_lock(cli_runs, tmp_path):
    lock = fno.acquire_lock(str(tmp_path), "NIFTY", "nifty_daily.py")
    try:
        assert nifty_bhav.main(["2026-10-12", "--folder", str(tmp_path)]) == 1
    finally:
        lock.release()
    assert cli_runs[0] == []


def test_cli_bhav_error_exit_1(monkeypatch, tmp_path):
    def fail(*a, **kw):
        raise bhav.BhavError("blocked")
    monkeypatch.setattr(nifty_bhav, "run_bhav", fail)
    assert nifty_bhav.main(["2026-10-12", "--folder", str(tmp_path)]) == 1


# --- dashboard routes ---------------------------------------------------------------------------

@pytest.fixture
def client(monkeypatch):
    monkeypatch.setattr(dashboard, "save_env", lambda updates: None)
    with JOB_LOCK:
        JOB["state"] = "idle"
    c = app.test_client()
    with c.session_transaction() as s:
        s["access_token"] = "x"
        s["user"] = {"user_id": "AB1234", "user_name": "Demo User", "email": "", "broker": "", "login_time": ""}
    return c


def checked_day(root, d, status="ok", counts=None, units=None, rows=()):
    """A saved Kite day with a stored bhavcopy check (status None = pending)."""
    day = root / "NIFTY_FnO_1min" / d
    day.mkdir(parents=True, exist_ok=True)
    (day / "summary.csv").write_text("x")
    if status:
        (day / "bhavcopy").mkdir(exist_ok=True)
        (day / "bhavcopy" / "bhav_check.json").write_text(json.dumps(
            {"status": status, "line": "Bhavcopy check: OK" if status == "ok" else "Bhavcopy check: CHECK (1 missing)",
             "counts": counts or {"missing": 0}, "units": units or {"volume": 65, "oi": 1}, "checked_at": "t"}))
        fno.write_csv(str(day / "bhavcopy" / "bhav_compare.csv"), bhav.COMPARE_COLUMNS,
                      [{c: r.get(c, "") for c in bhav.COMPARE_COLUMNS} for r in rows])
    return day


def test_list_days_latest_first(tmp_path):
    assert bhav.list_days(str(tmp_path / "nowhere")) == [] and not (tmp_path / "nowhere").exists()
    checked_day(tmp_path, "2026-10-07", "ok")
    checked_day(tmp_path, "2026-10-08", "diffs", {"missing": 1})
    checked_day(tmp_path, "2026-10-09", None)
    (tmp_path / "NIFTY_FnO_1min" / "2026-10-10").mkdir()          # no summary.csv: not a saved day
    days = bhav.list_days(str(tmp_path))
    assert [(d["date"], d["status"]) for d in days] == [("2026-10-09", "pending"), ("2026-10-08", "diffs"),
                                                        ("2026-10-07", "ok")]
    assert days[0]["line"] == bhav.PENDING and days[1]["counts"] == {"missing": 1}


def test_banners():
    today = date(2026, 10, 22)
    days = [{"date": f"2026-10-{d:02d}", "status": st} for d, st in
            [(22, "pending"), (21, "pending"), (20, "pending"), (19, "ok"), (16, "ok"), (15, "ok"), (14, "diffs"),
             (13, "ok"), (12, "ok"), (9, "ok"), (8, "ok"), (7, "diffs")]]   # 7 Oct is the 12th day: too old
    b = bhav.banners(days, today)
    assert b["check"] == ["2026-10-14"]
    assert b["pending_old"] == ["2026-10-20"]                    # today and yesterday may not be out yet


def test_diff_rows(tmp_path, monkeypatch):
    monkeypatch.setattr(fno, "now_ist", lambda: datetime(2026, 10, 14, 18, 0, tzinfo=IST))
    checked_day(tmp_path, DAY, "diffs", rows=[
        {"Contract": "NIFTY27OCT26", "Expiry": "2026-10-27", "Type": "FUT", "Bhav_volume": "10", "Kite_volume": "650", "Status": "OK"},
        {"Contract": "NIFTY2500013OCT26CE", "Expiry": "2026-10-13", "Type": "CE", "Bhav_volume": "3", "Status": "MISSING_IN_KITE"},
        {"Contract": "NIFTY2500020OCT26PE", "Expiry": "2026-10-20", "Type": "PE", "Bhav_volume": "2", "Status": "MISSING_IN_KITE"},
        {"Contract": "NIFTY2510020OCT26PE", "Expiry": "2026-10-20", "Type": "PE", "Bhav_volume": "2", "Kite_volume": "140",
         "Status": "VOLUME_DIFF+OI_DIFF"}])
    rows, check = bhav.diff_rows(str(tmp_path), "NIFTY", RUN_DATE)
    assert [r["Contract"] for r in rows] == ["NIFTY2500013OCT26CE", "NIFTY2500020OCT26PE", "NIFTY2510020OCT26PE"]
    assert [r["Live"] for r in rows] == [False, True, True]       # 13 Oct expired by 14 Oct
    assert rows[2]["Bhav_volume_shares"] == "130"
    with pytest.raises(bhav.BhavError):
        bhav.diff_rows(str(tmp_path), "NIFTY", date(2026, 10, 9))


def test_redownload_fixes_missing_contract(tmp_path, fake_kite):
    data = tmp_path / "data"
    run(data)
    day = data / "NIFTY_FnO_1min" / DAY
    os.remove(day / "options" / "NIFTY2510013OCT26CE.csv")        # lost: the bhavcopy says it traded
    status, line, _ = bhav.run_bhav(RUN_DATE, str(data), bhav_file=day_bhav(tmp_path / "b.zip"), log=lambda m: None)
    assert line == "Bhavcopy check: CHECK (2 missing); liquid buckets: 2 missing, 0 diffs"
    rows, _ = bhav.diff_rows(str(data), "NIFTY", RUN_DATE)
    assert {r["Contract"]: r["Live"] for r in rows}["NIFTY2510013OCT26CE"] is True
    fake_kite.calls.clear()
    saved, msg = fno.redownload_contract("key", ACCESS, RUN_DATE, str(data), "NIFTY2510013OCT26CE", speed=1000,
                                         log=lambda m: None)
    assert saved and "5 bars" in msg and [t for t, *_ in fake_kite.calls] == [11]
    assert (day / "options" / "NIFTY2510013OCT26CE.csv").exists()
    status, line, _ = bhav.check_saved(str(day), RUN_DATE, log=lambda m: None)
    assert line.startswith("Bhavcopy check: CHECK (1 missing);")  # only the contract Kite never traded is left
    db = open_db(str(data))
    assert db.execute("SELECT status FROM chunks WHERE run = ? AND symbol = 'NIFTY2610625100CE'",
                      (fno.run_key("NIFTY", RUN_DATE),)).fetchone()[0] == "DONE"
    db.close()


def test_redownload_refuses_expired_or_unknown(tmp_path, fake_kite, monkeypatch):
    data = tmp_path / "data"
    run(data)
    set_now(monkeypatch, date(2026, 10, 14))                     # the 13 Oct weekly has expired
    fake_kite.calls.clear()
    saved, msg = fno.redownload_contract("key", ACCESS, RUN_DATE, str(data), "NIFTY2510013OCT26CE", log=lambda m: None)
    assert not saved and "expired" in msg
    saved, msg = fno.redownload_contract("key", ACCESS, RUN_DATE, str(data), "NIFTY9999913OCT26CE", log=lambda m: None)
    assert not saved and "not in the instrument list" in msg
    assert fake_kite.calls == []


def test_route_bhav_status(client, tmp_path, monkeypatch):
    monkeypatch.setattr(fno, "now_ist", lambda: datetime(2026, 10, 12, 18, 0, tzinfo=IST))
    checked_day(tmp_path, "2026-10-08", "diffs")
    checked_day(tmp_path, "2026-10-09", None)
    j = client.get("/fno/bhav-status", query_string={"folder": str(tmp_path)}).get_json()
    assert [d["date"] for d in j["days"]] == ["2026-10-09", "2026-10-08"]
    assert j["banners"] == {"check": ["2026-10-08"], "minor": [], "pending_old": ["2026-10-09"]}


def test_route_bhav_diff(client, tmp_path):
    checked_day(tmp_path, DAY, "diffs", rows=[{"Contract": "A", "Expiry": "2030-01-01", "Status": "MISSING_IN_KITE"},
                                              {"Contract": "B", "Expiry": "2030-01-01", "Status": "OK"}])
    j = client.get("/fno/bhav-diff", query_string={"folder": str(tmp_path), "date": DAY}).get_json()
    assert [r["Contract"] for r in j["rows"]] == ["A"] and j["rows"][0]["Live"] is True
    j = client.get("/fno/bhav-diff", query_string={"folder": str(tmp_path), "date": "2026-10-09"}).get_json()
    assert "not been checked" in j["error"]


def test_route_check_date_and_missing(client, tmp_path, monkeypatch):
    calls = []

    def fake(d, folder, underlying, bhav_file=None, log=print, source="kite"):
        assert source == "kite"
        calls.append(d)
        if d == date(2026, 10, 9):
            return "not_available", bhav.PENDING, None
        return "diffs", "Bhavcopy check: CHECK (1 missing)", None
    monkeypatch.setattr(bhav, "run_bhav", fake)
    j = client.post("/fno/bhav-check", data={"folder": str(tmp_path), "date": DAY}).get_json()
    assert j == {"results": [{"date": DAY, "status": "diffs", "line": "Bhavcopy check: CHECK (1 missing)"}]}
    checked_day(tmp_path, "2026-10-08", None)
    checked_day(tmp_path, "2026-10-09", None)
    checked_day(tmp_path, "2026-10-07", "ok")
    calls.clear()
    j = client.post("/fno/bhav-check", data={"folder": str(tmp_path), "missing": "1"}).get_json()
    assert [r["status"] for r in j["results"]] == ["diffs", "not_available"]
    assert calls == [date(2026, 10, 8), date(2026, 10, 9)]


def test_route_check_refuses_when_locked_or_running(client, tmp_path, monkeypatch):
    monkeypatch.setattr(bhav, "run_bhav", lambda *a, **kw: pytest.fail("ran"))
    monkeypatch.setattr(dashboard, "load_env", lambda: {"KITE_ACCESS_TOKEN": "x"})   # never the real .env
    lock = fno.acquire_lock(str(tmp_path), "NIFTY", "nifty_daily.py")
    try:
        j = client.post("/fno/bhav-check", data={"folder": str(tmp_path), "date": DAY}).get_json()
        assert "nifty_daily.py" in j["error"]
    finally:
        lock.release()
    with JOB_LOCK:
        JOB["state"] = "running"
    try:
        assert "running" in client.post("/fno/bhav-check", data={"folder": str(tmp_path), "date": DAY}).get_json()["error"]
        j = client.post("/fno/redownload", data={"folder": str(tmp_path), "date": DAY, "contract": "X"}).get_json()
        assert "running" in j["error"]
    finally:
        with JOB_LOCK:
            JOB["state"] = "idle"


def test_route_redownload(client, tmp_path, monkeypatch):
    monkeypatch.setattr(dashboard, "load_env", lambda: {"KITE_ACCESS_TOKEN": "x"})
    monkeypatch.setattr(dashboard, "credentials", lambda: ("key", ""))
    monkeypatch.setattr(fno, "redownload_contract", lambda *a, **kw: (True, "NIFTY2500013OCT26CE: saved 5 bars."))
    monkeypatch.setattr(bhav, "check_saved", lambda *a, **kw: ("ok", "Bhavcopy check: OK", None))
    form = {"folder": str(tmp_path), "date": DAY, "contract": "NIFTY2500013OCT26CE"}
    assert client.post("/fno/redownload", data=form).get_json() == \
        {"saved": True, "message": "NIFTY2500013OCT26CE: saved 5 bars.", "status": "ok", "line": "Bhavcopy check: OK"}
    monkeypatch.setattr(fno, "redownload_contract", lambda *a, **kw: (False, "NIFTY2500013OCT26CE expired on 2026-10-13"))
    j = client.post("/fno/redownload", data=form).get_json()
    assert j["saved"] is False and "expired" in j["message"]


def test_old_routes_removed(client):
    assert client.post("/download/fno-bhav").status_code == 404
    assert client.get("/download/fno-bhav-status").status_code == 404


# --- buckets: impact, not just counts -----------------------------------------------------------

def test_parse_name():
    assert bhav.parse_name("NIFTY2500013OCT26CE") == (date(2026, 10, 13), "CE", 25000.0)
    assert bhav.parse_name("NIFTY27OCT26") == (date(2026, 10, 27), "FUT", "")
    assert bhav.parse_name("NIFTY2500013OCT26") is None and bhav.parse_name("NIFTY13OCT26CE") is None
    assert bhav.parse_name("SENSEX29OCT26", "SENSEX") == (date(2026, 10, 29), "FUT", "")


def by_contract(rows):
    return {r["Contract"]: r for r in rows}


def test_buckets_and_liquid_counts(tmp_path):
    d = str(tmp_path)
    # option expiries 13, 20, 27 Oct (27 Oct is also the month future), 29 Dec far; futures 27 Oct, 24 Nov, 29 Dec
    kite_file(d, "NIFTY27OCT26", [100.0], vol=6500, oi_last=1000)
    kite_file(d, "NIFTY24NOV26", [101.0], vol=6500, oi_last=1000)
    kite_file(d, "NIFTY2500013OCT26CE", [10.0], vol=66950, oi_last=1000)   # 3% / 30 lots volume diff, current week
    kite_file(d, "NIFTY2500020OCT26CE", [10.0], vol=650, oi_last=1000)
    kite_file(d, "NIFTY3000029DEC26CE", [1.0], vol=65, oi_last=10)         # extra, absent from bhavcopy: far
    items = bhav.parse_bhav([
        row("IDF", "2026-10-27", vol=100, oi=1000, last=100, lot=65),
        row("IDF", "2026-11-24", vol=100, oi=1000, last=101, lot=65),
        row("IDF", "2026-12-29", vol=0, lot=65),
        row("IDO", "2026-10-13", "25000.00", "CE", vol=1000, oi=1000, last=10, lot=65),
        row("IDO", "2026-10-20", "25000.00", "CE", vol=10, oi=1000, last=10, lot=65),
        row("IDO", "2026-10-27", "25000.00", "PE", vol=0, lot=65),
        row("IDO", "2026-12-29", "12000.00", "PE", vol=2, last=1, lot=65),     # traded, no file: far missing
    ])
    rows, counts, _ = bhav.compare(items, d)
    r = by_contract(rows)
    assert r["NIFTY2500013OCT26CE"]["Bucket"] == "Current week"
    assert r["NIFTY2500020OCT26CE"]["Bucket"] == "Next week"
    assert r["NIFTY2500027OCT26PE"]["Bucket"] == "Current month"
    assert r["NIFTY27OCT26"]["Bucket"] == "Current month" and r["NIFTY24NOV26"]["Bucket"] == "Next month"
    assert r["NIFTY29DEC26"]["Bucket"] == "Far" and r["NIFTY1200029DEC26PE"]["Bucket"] == "Far"
    extra = r["NIFTY3000029DEC26CE"]
    assert (extra["Expiry"], extra["Bucket"], extra["Type"], extra["Strike"], extra["Status"]) == \
        ("2026-12-29", "Far", "CE", "30000", "EXTRA_IN_KITE")
    # nothing hidden: totals keep every difference, the liquid headline sits beside them
    assert (counts["missing"], counts["extra"], counts["volume_diffs"]) == (1, 1, 1)
    assert (counts["liquid_missing"], counts["liquid_diffs"], counts["liquid_volume_diffs"], counts["far_problems"]) == (0, 1, 1, 2)
    assert bhav.check_status(counts) == ("Bhavcopy check: CHECK (1 missing, 1 extra, 1 volume diffs); "
                                         "liquid buckets: 0 missing, 1 diffs")
    assert bhav.severity(counts) == "major"      # a volume diff with Kite volume, and a far missing contract


def test_liquid_oi_and_price_diffs_are_minor(tmp_path):
    d = str(tmp_path)
    kite_file(d, "NIFTY27OCT26", [100.0], vol=650, oi_last=1100)            # OI +10%
    kite_file(d, "NIFTY2500013OCT26CE", [12.0], vol=650, oi_last=1000)      # price 12 vs 10
    items = bhav.parse_bhav([row("IDF", "2026-10-27", vol=10, oi=1000, last=100, lot=65),
                             row("IDO", "2026-10-13", "25000.00", "CE", vol=10, oi=1000, last=10, lot=65)])
    rows, counts, _ = bhav.compare(items, d)
    assert set(statuses(rows).values()) == {"OI_DIFF", "PRICE_DIFF"}
    assert (counts["liquid_diffs"], counts["liquid_missing"], counts["liquid_volume_diffs"]) == (2, 0, 0)
    assert bhav.severity(counts) == "minor"


def test_monthly_option_expiry_gets_two_labels(tmp_path):
    items = bhav.parse_bhav([row("IDF", "2026-10-27", vol=1), row("IDO", "2026-10-20", "25000.00", "CE", vol=1),
                             row("IDO", "2026-10-27", "25000.00", "CE", vol=1)])
    r = by_contract(bhav.compare(items, str(tmp_path))[0])
    assert r["NIFTY2500027OCT26CE"]["Bucket"] == "Next week + Current month"


def test_severity():
    base = {"missing": 0, "extra": 0, "zero_volume": 0, "volume_diffs": 0, "oi_diffs": 0, "price_diffs": 0,
            "liquid_missing": 0, "liquid_diffs": 0, "liquid_volume_diffs": 0, "far_problems": 0, "red_volume_diffs": 0}
    # known Kite limitations are amber even in liquid buckets
    assert bhav.severity({**base, "extra": 2, "oi_diffs": 3, "price_diffs": 1, "liquid_diffs": 6}) == "minor"
    assert bhav.severity({**base, "zero_volume": 1, "liquid_diffs": 1}) == "minor"
    assert bhav.severity({**base, "zero_volume": 1, "volume_diffs": 1, "liquid_volume_diffs": 1}) == "minor"
    # red: missing in any bucket, or a volume diff where Kite has volume
    assert bhav.severity({**base, "missing": 1, "far_problems": 1}) == "major"
    assert bhav.severity({**base, "volume_diffs": 1, "red_volume_diffs": 1, "far_problems": 1}) == "major"
    assert bhav.severity({"missing": 0, "volume_diffs": 1}) == "major"         # checked before red_volume_diffs
    assert bhav.severity({"missing": 2}) == "major"
    assert bhav.severity(base, "ok") == ""


def test_list_days_severity_and_banners(tmp_path):
    checked_day(tmp_path, "2026-10-06", "diffs")                           # old check: no severity -> major
    checked_day(tmp_path, "2026-10-07", "diffs")
    checked_day(tmp_path, "2026-10-08", "diffs", counts={"missing": 0, "volume_diffs": 0, "oi_diffs": 2, "liquid_missing": 0,
                                                         "liquid_diffs": 2, "liquid_volume_diffs": 0, "far_problems": 0})
    path = tmp_path / "NIFTY_FnO_1min" / "2026-10-07" / "bhavcopy" / "bhav_check.json"
    data = json.loads(path.read_text())
    data["severity"] = "far"                                               # name used by the previous version
    path.write_text(json.dumps(data))
    checked_day(tmp_path, "2026-10-09", "ok")
    days = bhav.list_days(str(tmp_path))
    assert [(d["date"], d["severity"]) for d in days] == [("2026-10-09", ""), ("2026-10-08", "minor"),
                                                          ("2026-10-07", "minor"), ("2026-10-06", "major")]
    assert bhav.banners(days, date(2026, 10, 12)) == {"check": ["2026-10-06"], "minor": ["2026-10-08", "2026-10-07"],
                                                      "pending_old": []}


def test_diff_rows_sorted_by_bhav_volume(tmp_path):
    checked_day(tmp_path, DAY, "diffs", rows=[
        {"Contract": "A", "Expiry": "2030-01-01", "Bucket": "Far", "Bhav_volume": "2", "Status": "MISSING_IN_KITE"},
        {"Contract": "B", "Expiry": "2030-01-01", "Bucket": "Current week", "Bhav_volume": "50", "Status": "VOLUME_DIFF"},
        {"Contract": "C", "Expiry": "2030-01-01", "Bucket": "Far", "Bhav_volume": "", "Kite_volume": "65", "Status": "EXTRA_IN_KITE"},
        {"Contract": "D", "Expiry": "2030-01-01", "Bucket": "Next week", "Bhav_volume": "9", "Status": "OI_DIFF"},
        {"Contract": "E", "Expiry": "2030-01-01", "Bucket": "Far", "Bhav_volume": "99", "Status": "OK"}])
    rows, _ = bhav.diff_rows(str(tmp_path), "NIFTY", RUN_DATE)
    assert [(r["Contract"], r["Liquid"]) for r in rows] == [("B", True), ("D", True), ("A", False), ("C", False)]
    assert rows[0]["Bhav_volume_shares"] == "3250"


def test_diff_rows_fills_bucket_for_old_checks(tmp_path, fake_kite):
    data = tmp_path / "data"
    run(data)
    bhav.run_bhav(RUN_DATE, str(data), bhav_file=day_bhav(tmp_path / "b.zip"), log=lambda m: None)
    cmp = data / "NIFTY_FnO_1min" / DAY / "bhavcopy" / "bhav_compare.csv"
    old = [{k: v for k, v in r.items() if k != "Bucket"} for r in csv.DictReader(open(cmp, newline=""))]
    fno.write_csv(str(cmp), [c for c in bhav.COMPARE_COLUMNS if c != "Bucket"], old)   # as written before buckets
    rows, _ = bhav.diff_rows(str(data), "NIFTY", RUN_DATE)
    assert [(r["Contract"], r["Bucket"]) for r in rows] == [("NIFTY2515013OCT26CE", "Current week")]


def test_load_eod_for_backtests(tmp_path):
    from backfill.fno_loader import load_eod
    src = tmp_path / "b.zip"
    src.write_bytes(zip_bytes([row("IDF", "2026-10-27", vol=3, oi=130000, last=22625, close=22612.6, settle=22612.6),
                               row("IDO", "2026-10-13", "25000.00", "CE", vol=2, oi=6500, last=10, close=10.5, settle=10.45)]))
    bhav.run_bhav(RUN_DATE, str(tmp_path / "data"), bhav_file=str(src), log=lambda m: None)
    eod = load_eod(str(tmp_path / "data"), DAY).set_index("Contract")
    assert list(eod.columns) == ["Expiry", "Type", "Strike", "Close", "Last", "Settle", "OI", "Volume_contracts", "Lot"]
    assert eod.loc["NIFTY27OCT26", "Settle"] == 22612.6 and eod.loc["NIFTY27OCT26", "OI"] == 130000
    assert eod.loc["NIFTY2500013OCT26CE", "Close"] == 10.5 and eod.loc["NIFTY2500013OCT26CE", "Strike"] == 25000
    with pytest.raises(FileNotFoundError):
        load_eod(str(tmp_path / "data"), "2026-10-09")


# --- zero-volume Kite files ---------------------------------------------------------------------

def zero_file(day_dir, name, closes):
    candles = [{"date": datetime(2026, 10, 12, 13, 12 + k, tzinfo=IST), "open": c, "high": c, "low": c, "close": c,
                "volume": 0, "oi": 1625} for k, c in enumerate(closes)]
    fno.write_contract(os.path.join(day_dir, "options", name + ".csv"), name, candles, keep_zero=True)


def test_zero_volume_file_with_bhav_trade(tmp_path):
    d = str(tmp_path)
    kite_file(d, "NIFTY27OCT26", [100.0], vol=650)
    zero_file(d, "NIFTY2415013OCT26PE", [1571.0, 1571.0])
    items = bhav.parse_bhav([row("IDF", "2026-10-27", vol=10, last=100, lot=65),
                             row("IDO", "2026-10-13", "24150.00", "PE", vol=1, last=1571, oi=1625, lot=65)])
    rows, counts, _ = bhav.compare(items, d)
    assert statuses(rows)["NIFTY2415013OCT26PE"] == "ZERO_VOLUME_IN_KITE"   # one lot: under the volume limits
    assert counts["zero_volume"] == 1 and counts["missing"] == 0
    assert counts["liquid_missing"] == 0 and counts["liquid_volume_diffs"] == 0
    assert bhav.severity(counts) == "minor"
    assert "1 zero-volume in Kite" in bhav.check_status(counts)


def test_zero_volume_file_over_limits_is_also_volume_diff(tmp_path):
    d = str(tmp_path)
    kite_file(d, "NIFTY27OCT26", [100.0], vol=650)
    zero_file(d, "NIFTY2415013OCT26PE", [1571.0])
    items = bhav.parse_bhav([row("IDF", "2026-10-27", vol=10, last=100, lot=65),
                             row("IDO", "2026-10-13", "24150.00", "PE", vol=30, last=1571, oi=1625, lot=65)])
    rows, counts, _ = bhav.compare(items, d)
    assert statuses(rows)["NIFTY2415013OCT26PE"] == "ZERO_VOLUME_IN_KITE+VOLUME_DIFF"
    # Kite volume 0: a known limitation, amber even in a liquid bucket; the difference is still counted
    assert counts["volume_diffs"] == 1 and counts["red_volume_diffs"] == 0 and bhav.severity(counts) == "minor"


def test_volume_diff_with_kite_volume_is_red_in_far_bucket(tmp_path):
    d = str(tmp_path)
    kite_file(d, "NIFTY27OCT26", [100.0], vol=650)
    kite_file(d, "NIFTY3000028DEC27PE", [5000.0], vol=65)     # 1 lot in Kite, 30 in the bhavcopy
    items = bhav.parse_bhav([row("IDF", "2026-10-27", vol=10, last=100, lot=65),
                             row("IDO", "2026-10-13", "25000.00", "CE", lot=65),   # nearer expiries, untraded
                             row("IDO", "2026-10-20", "25000.00", "CE", lot=65),
                             row("IDO", "2027-12-28", "30000.00", "PE", vol=30, last=5000, oi=1000, lot=65)])
    rows, counts, _ = bhav.compare(items, d)
    assert by_contract(rows)["NIFTY3000028DEC27PE"]["Bucket"] == "Far"
    assert counts["red_volume_diffs"] == 1 and bhav.severity(counts) == "major"


def test_kept_zero_volume_file_without_bhav_trade_is_extra(tmp_path):
    d = str(tmp_path)
    kite_file(d, "NIFTY27OCT26", [100.0], vol=650)
    zero_file(d, "NIFTY2415013OCT26PE", [1571.0])
    zero_file(d, "NIFTY2420013OCT26PE", [1600.0])   # not in the bhavcopy at all
    items = bhav.parse_bhav([row("IDF", "2026-10-27", vol=10, last=100, lot=65),
                             row("IDO", "2026-10-13", "24150.00", "PE", vol=0, lot=65)])
    rows, counts, _ = bhav.compare(items, d)
    st = statuses(rows)
    assert st["NIFTY2415013OCT26PE"] == "EXTRA_IN_KITE" and st["NIFTY2420013OCT26PE"] == "EXTRA_IN_KITE"
    assert counts["zero_volume"] == 0 and counts["extra"] == 2


# --- downloads of the check (backfill/bhav_export.py, /fno/bhav-export, nifty_bhav.py --export) ---

EXPORT_ROWS = [
    {"Contract": "A", "Expiry": "2030-01-01", "Bucket": "Far", "Bhav_volume": "2", "Status": "MISSING_IN_KITE"},
    {"Contract": "B", "Expiry": "2030-01-01", "Bucket": "Current week", "Bhav_volume": "50", "Kite_volume": "1625",
     "Bhav_OI": "2036458500", "Kite_OI": "2036458000", "Bhav_last": "1571.5", "Kite_last": "1570",
     "Kite_last_regular": "1569.95", "Last_bar": "15:31:00", "Status": "VOLUME_DIFF+PRICE_DIFF",
     "Note": "trades after Kite's last bar"},
    {"Contract": "D", "Expiry": "2030-01-01", "Bucket": "Next week", "Bhav_volume": "9", "Status": "OI_DIFF"},
    {"Contract": "F", "Expiry": "2030-01-01", "Bucket": "Current week", "Bhav_volume": "1", "Status": "MISSING_IN_KITE"},
    {"Contract": "E", "Expiry": "2030-01-01", "Bucket": "Current week", "Bhav_volume": "99", "Status": "OK"}]


def exported(c, folder, **params):
    return c.get("/fno/bhav-export", query_string={"folder": str(folder), "underlying": "NIFTY", **params})


def csv_dicts(data):
    return list(csv.DictReader(io.StringIO(data.decode("utf-8"))))


def checked_fake_day(tmp_path):
    """The fake Kite day of test_fno checked against day_bhav(): a real bhav_compare.csv and summary.csv."""
    data = tmp_path / "data"
    run(data)
    bhav.run_bhav(RUN_DATE, str(data), bhav_file=day_bhav(tmp_path / "b.zip"), log=lambda m: None)
    return data


def tree_bytes(path):
    return {os.path.relpath(os.path.join(d, f), path): open(os.path.join(d, f), "rb").read()
            for d, _, fs in os.walk(path) for f in fs}


def test_export_view_keeps_bucket_filter_and_order(tmp_path, client):
    checked_day(tmp_path, DAY, "diffs", rows=EXPORT_ROWS)
    r = exported(client, tmp_path, date=DAY, scope="view", bucket="Current week")
    assert r.status_code == 200
    assert r.headers["Content-Disposition"] == "attachment; filename=bhavcopy_check_2026-10-12_current_week.csv"
    rows = csv_dicts(r.data)
    assert [x["Contract"] for x in rows] == ["B", "F"]          # not OK only, biggest bhavcopy volume first
    assert list(rows[0]) == bhav_export.CONTRACT_COLUMNS
    b = rows[0]
    assert (b["Bhav_volume"], b["Bhav_volume_shares"], b["Kite_volume"], b["Kite_regular_close"], b["Kite_last_bar_time"],
            b["Status_note"], b["Expiry"], b["Bucket"]) == \
        ("50", "3250", "1625", "1569.95", "15:31:00", "trades after Kite's last bar", "2030-01-01", "Current week")
    assert [x["Contract"] for x in csv_dicts(exported(client, tmp_path, date=DAY, scope="view", bucket="liquid").data)] == \
        ["B", "D", "F"]
    r = exported(client, tmp_path, date=DAY, scope="view", bucket="all")
    assert r.headers["Content-Disposition"].endswith("bhavcopy_check_2026-10-12_all.csv")
    assert [x["Contract"] for x in csv_dicts(r.data)] == ["B", "D", "A", "F"]


def test_export_numbers_are_plain(tmp_path, client):
    checked_day(tmp_path, DAY, "diffs", rows=EXPORT_ROWS)
    data = exported(client, tmp_path, date=DAY, scope="all").data
    text = data.decode("utf-8")
    assert text.startswith("Contract,Expiry,Bucket,")           # header first, no comment lines
    for r in csv_dicts(data):
        assert not any("," in v for v in r.values() if v != r["Status_note"])
    df = pd.read_csv(io.BytesIO(data))
    assert len(df) == 5 and df["Bhav_OI"].dropna().iloc[0] == 2036458500
    assert str(df["Bhav_last"].dtype) == "float64" and str(df["Bhav_volume"].dtype) == "int64"


def test_export_all_contracts_includes_ok_rows(tmp_path, client, fake_kite):
    data = checked_fake_day(tmp_path)
    r = exported(client, data, date=DAY, scope="all")
    assert r.headers["Content-Disposition"] == "attachment; filename=bhavcopy_check_2026-10-12_all_contracts.csv"
    rows = {x["Contract"]: x for x in csv_dicts(r.data)}
    assert rows["NIFTY27OCT26"]["Status"] == "OK" and rows["NIFTY2515013OCT26CE"]["Status"] == "MISSING_IN_KITE"
    assert rows["NIFTY27OCT26"]["Bucket"] == "Current month" and rows["NIFTY27OCT26"]["Expiry"] == "2026-10-27"
    assert rows["NIFTY27OCT26"]["Bhav_volume_shares"] == "500"


def test_export_xlsx_sheets(tmp_path, client, fake_kite):
    from openpyxl import load_workbook
    data = checked_fake_day(tmp_path)
    r = exported(client, data, date=DAY, scope="xlsx")
    assert r.status_code == 200 and r.headers["Content-Disposition"] == "attachment; filename=bhavcopy_check_2026-10-12.xlsx"
    wb = load_workbook(io.BytesIO(r.data))
    assert wb.sheetnames == ["Summary", "Not OK", "All contracts", "Settlement"]
    for ws in wb.worksheets:
        assert ws.freeze_panes and ws.auto_filter.ref
    summary = wb["Summary"]
    assert [c.value for c in summary[1]] == bhav_export.DAY_COLUMNS
    assert summary["A2"].value == DAY and summary["B2"].value == "CHECK" and summary["M2"].value == "OFF"
    assert [c.value for c in summary[4] if c.value is not None] == bhav_export.EXPIRY_COLUMNS and summary.freeze_panes == "A5"
    week = next(row for row in summary.iter_rows(min_row=5, values_only=True) if row[1] == "Current week")
    assert week[0] == "2026-10-13" and week[4:8] == (3, 2, 1, 0)   # traded, files, missing, extra as numbers
    assert wb["Not OK"].freeze_panes == "A2"
    assert [r[0] for r in wb["Not OK"].iter_rows(min_row=2, values_only=True)] == ["NIFTY2515013OCT26CE"]
    settle = wb["Settlement"]
    assert [c.value for c in settle[1]] == bhav_export.SETTLEMENT_COLUMNS
    fut = next(r for r in settle.iter_rows(min_row=2, values_only=True) if r[0] == "NIFTY27OCT26")
    assert fut[4] == 25130 and isinstance(fut[6], int)


def test_export_days_summary(tmp_path, client):
    checked_day(tmp_path, "2026-10-08", "diffs", counts={"missing": 2, "volume_diffs": 0, "liquid_missing": 1,
                                                         "liquid_diffs": 0, "liquid_volume_diffs": 0, "zero_volume": 3})
    checked_day(tmp_path, "2026-10-09", "ok", counts={"missing": 0, "bhav_traded": 900})
    checked_day(tmp_path, "2026-10-12", None)   # pending: left out
    r = exported(client, tmp_path, scope="days")
    assert r.headers["Content-Disposition"].startswith("attachment; filename=bhavcopy_check_all_days_")
    rows = csv_dicts(r.data)
    assert list(rows[0]) == bhav_export.DAY_COLUMNS
    assert [(x["Date"], x["Status"], x["Missing"], x["Zero_volume_in_kite"], x["Bhav_traded"]) for x in rows] == [
        ("2026-10-09", "OK", "0", "", "900"), ("2026-10-08", "CHECK", "2", "3", "")]


def test_check_saves_checked_days_summary(tmp_path, fake_kite):
    data = checked_fake_day(tmp_path)
    out = data / "NIFTY_FnO_1min" / DAY / "bhavcopy" / bhav.DAYS_SUMMARY_FILE
    assert out.read_bytes() == bhav_export.export(str(data), "NIFTY", "days")[1]   # only one day checked
    assert read_csv(out)[0] == bhav_export.DAY_COLUMNS
    # with another checked day, the day's file still holds only its own row; the download lists both
    other = checked_day(data, "2026-10-09", "ok", counts={"missing": 0, "bhav_traded": 900})
    bhav.check_saved(str(data / "NIFTY_FnO_1min" / DAY), RUN_DATE, log=lambda m: None)
    assert [r["Date"] for r in dict_rows(out)] == [DAY]
    assert len(bhav_export.days_rows(str(data), "NIFTY")) == 2
    assert not (other / "bhavcopy" / bhav.DAYS_SUMMARY_FILE).exists()


def test_checked_days_summary_needs_a_saved_day(tmp_path):
    day = tmp_path / "NIFTY_FnO_1min" / DAY      # contract files and a bhavcopy, but no summary.csv
    kite_file(str(day), "NIFTY27OCT26", [100.0])
    (day / "bhavcopy").mkdir()
    fno.write_csv(str(day / "bhavcopy" / "nifty_bhav.csv"), bhav.BHAV_HEADER, [row("IDF", "2026-10-27", vol=1, last=100)])
    bhav.check_saved(str(day), RUN_DATE, log=lambda m: None)
    assert (day / "bhavcopy" / "bhav_check.json").exists()
    assert not (day / "bhavcopy" / bhav.DAYS_SUMMARY_FILE).exists()


def test_export_unchecked_day_is_404(tmp_path, client, capsys):
    checked_day(tmp_path, DAY, None)
    for scope in ("view", "all", "xlsx"):
        r = exported(client, tmp_path, date=DAY, scope=scope)
        assert r.status_code == 404 and "has not been checked" in r.get_json()["error"]
    r = exported(client, tmp_path, scope="days")
    assert r.status_code == 404 and "No day" in r.get_json()["error"]
    assert exported(client, tmp_path, date=DAY, scope="nope").status_code == 400
    assert nifty_bhav.main(["--export", DAY, "--folder", str(tmp_path)]) == 1
    assert "Not exported" in capsys.readouterr().err
    assert not (tmp_path / "NIFTY_FnO_1min" / DAY / "bhavcopy").exists()


def test_export_changes_nothing_in_day_folder(tmp_path, client, fake_kite):
    data = checked_fake_day(tmp_path)
    before = tree_bytes(data)
    for params in ({"scope": "days"}, {"date": DAY, "scope": "view", "bucket": "Current week"},
                   {"date": DAY, "scope": "all"}, {"date": DAY, "scope": "xlsx"}):
        assert exported(client, data, **params).status_code == 200
    assert tree_bytes(data) == before


def test_cli_export(tmp_path, capsys, fake_kite):
    data = checked_fake_day(tmp_path)
    exports = data / "NIFTY_FnO_1min" / DAY / "bhavcopy" / "exports"
    assert nifty_bhav.main(["--export", DAY, "--folder", str(data)]) == 0
    assert capsys.readouterr().out.strip() == str(exports / "bhavcopy_check_2026-10-12.xlsx")
    assert nifty_bhav.main(["--export", DAY, "--csv", "--folder", str(data)]) == 0
    assert (exports / "bhavcopy_check_2026-10-12_all_contracts.csv").exists()
    assert nifty_bhav.main(["--export-days", "--folder", str(data)]) == 0
    out = capsys.readouterr().out.strip().splitlines()[-1]
    assert os.path.dirname(out) == str(data / "NIFTY_FnO_1min") and os.path.basename(out).startswith("bhavcopy_check_all_days_")
    with pytest.raises(SystemExit):
        nifty_bhav.main(["--export-days", "--csv", "--folder", str(data)])
